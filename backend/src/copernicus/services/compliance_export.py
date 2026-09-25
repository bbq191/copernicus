"""合规审核报告导出（Excel）。"""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from copernicus.schemas.compliance import ComplianceResponse, Violation

_SEVERITY_LABEL = {"high": "高", "medium": "中", "low": "低"}
_STATUS_LABEL = {"pending": "待审", "confirmed": "已确认", "rejected": "已驳回"}
_SOURCE_LABEL = {"transcript": "语音转写", "ocr": "屏幕文字", "vision": "视觉"}

_DETAIL_HEADERS = [
    ("序号", 6), ("时间", 10), ("说话人", 12), ("风险", 6), ("规则编号", 9),
    ("规则内容", 30), ("原文", 44), ("违规原因", 44), ("置信度", 8),
    ("来源", 10), ("复核状态", 10), ("复核时间(UTC)", 22), ("复核备注", 30),
    ("补充证据(OCR)", 36),
]


def _force_text(ws) -> None:
    """报告内容来自转写与模型输出，属不可信文本：以 = 开头的字符串会被 openpyxl 存成公式，
    这里强制改回文本类型，防止用户打开报告时执行公式（公式注入）。"""
    for row in ws.iter_rows():
        for cell in row:
            if cell.data_type == "f":
                cell.data_type = "s"


def _violation_row(index: int, v: Violation) -> list[object]:
    return [
        index,
        v.timestamp,
        v.speaker,
        _SEVERITY_LABEL.get(v.severity, v.severity),
        v.rule_id,
        v.rule_content,
        v.original_text,
        v.reason,
        round(v.confidence, 2),
        _SOURCE_LABEL.get(v.source, v.source),
        _STATUS_LABEL.get(v.status, v.status),
        v.reviewed_at or "",
        v.review_note or "",
        v.evidence_text or "",
    ]


def _summary_rows(resp: ComplianceResponse) -> list[tuple[str, object]]:
    report = resp.report
    counts = {
        label: sum(1 for v in report.violations if v.status == key)
        for key, label in _STATUS_LABEL.items()
    }
    rows: list[tuple[str, object]] = [
        ("合规评分", report.compliance_score),
        ("审核规则数", report.total_rules),
        ("已检查段落", report.total_segments_checked),
        ("违规条目", len(report.violations)),
        *[(f"其中{label}", n) for label, n in counts.items()],
    ]
    if report.truncated:
        rows.append(
            ("完整性提示", f"文本超过审核上限，仅检查了前 {report.total_segments_checked} / {report.total_segments} 个段落")
        )
    if report.failed_chunks:
        rows.append(
            ("完整性提示", f"{report.failed_chunks} / {report.total_chunks} 个审核分块失败被跳过，可能存在漏检")
        )
    rows.append(("审核摘要", report.summary))
    return rows


def build_compliance_xlsx(resp: ComplianceResponse) -> bytes:
    wb = Workbook()

    overview = wb.active
    overview.title = "概览"
    for label, value in _summary_rows(resp):
        overview.append([label, value])
    for row in overview.iter_rows():
        row[0].font = Font(bold=True)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    overview.column_dimensions["A"].width = 14
    overview.column_dimensions["B"].width = 80

    detail = wb.create_sheet("违规明细")
    detail.append([h for h, _ in _DETAIL_HEADERS])
    for cell in detail[1]:
        cell.font = Font(bold=True)
    for i, (_, width) in enumerate(_DETAIL_HEADERS, start=1):
        detail.column_dimensions[get_column_letter(i)].width = width
    for n, v in enumerate(resp.report.violations, start=1):
        detail.append(_violation_row(n, v))
    for row in detail.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    detail.freeze_panes = "A2"

    _force_text(overview)
    _force_text(detail)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
