"""合规规则文件（CSV/XLSX）解析。

与 ComplianceService 的 LLM 审核编排是两个正交职责：这里只管"把上传的文件变成
结构化规则列表"，不依赖 LLM 客户端或审核流程，便于独立测试。
"""

import csv
import io
import re
from collections.abc import Iterable

from copernicus.exceptions import ComplianceError
from copernicus.schemas.compliance import ComplianceRule

_HEADER_KEYWORDS = ("必备要素", "检查", "标准", "序号", "注：")
_SKIP_CELLS = {"合格", "不涉及", "None", ""}


def parse_rule_file(file_bytes: bytes, filename: str) -> tuple[list[ComplianceRule], list[str]]:
    """解析规则文件，返回 (规则列表, few-shot 案例列表)。

    支持 CSV 和 XLSX 格式。
    CSV A 列为规则，B-G 列为实际检查结果（作为 few-shot）。
    """
    if filename.lower().endswith((".xlsx", ".xls")):
        return _parse_xlsx(file_bytes)
    return _parse_csv(file_bytes)


def _parse_rule_rows(
    rows: Iterable[list[str]],
) -> tuple[list[ComplianceRule], list[str]]:
    """统一的规则行解析逻辑，CSV 和 XLSX 共用。"""
    rules: list[ComplianceRule] = []
    examples: list[str] = []

    for cells in rows:
        col_a = cells[0].strip() if cells else ""
        if not col_a:
            continue
        if any(kw in col_a for kw in _HEADER_KEYWORDS):
            continue
        if col_a.startswith("存在的问题"):
            break

        rule_id, content = _split_rule_id(col_a, len(rules) + 1)
        if not content:
            continue

        rules.append(ComplianceRule(id=rule_id, content=content))

        for cell in cells[1:]:
            cell = cell.strip()
            if cell and cell not in _SKIP_CELLS:
                examples.append(f"规则{rule_id}({content[:20]}...): {cell}")

    return rules, examples


def _parse_csv(file_bytes: bytes) -> tuple[list[ComplianceRule], list[str]]:
    """解析 CSV 文件，返回 (规则列表, few-shot 案例列表)。"""
    text = _decode_bytes(file_bytes)
    rows = (row for row in csv.reader(io.StringIO(text)) if row)
    return _parse_rule_rows(rows)


def _parse_xlsx(file_bytes: bytes) -> tuple[list[ComplianceRule], list[str]]:
    """解析 XLSX 文件。"""
    try:
        import openpyxl
    except ImportError as e:
        raise ComplianceError("解析 XLSX 需要 openpyxl 库") from e

    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    if not wb.sheetnames:
        raise ComplianceError("XLSX 文件没有工作表")
    ws = wb[wb.sheetnames[0]]

    rows = (
        [str(c).strip() if c is not None else "" for c in row]
        for row in ws.iter_rows(values_only=True)
    )
    return _parse_rule_rows(rows)


def _decode_bytes(data: bytes) -> str:
    """尝试多种编码解码 CSV。"""
    for encoding in ("utf-8-sig", "utf-8", "gbk", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ComplianceError("无法解码文件，请确保编码为 UTF-8 或 GBK")


def _split_rule_id(text: str, fallback_id: int) -> tuple[int, str]:
    """从 '4全程双录：...' 或 '10资料归档...' 中分离编号和内容。"""
    match = re.match(r"^(\d+)\s*(.+)", text, re.DOTALL)
    if match:
        return int(match.group(1)), match.group(2).strip()
    return fallback_id, text.strip()
