from unittest.mock import MagicMock

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from copernicus.routers.compliance import router as compliance_router
from copernicus.schemas.compliance import (
    ComplianceReport,
    ComplianceResponse,
    ComplianceRule,
    Violation,
)


def _violation(rule_id: int, ts: int, **kw) -> Violation:
    return Violation(
        rule_id=rule_id,
        rule_content="规则",
        reason="原因",
        confidence=0.9,
        timestamp_ms=ts,
        **kw,
    )


def _response(violations: list[Violation]) -> dict:
    resp = ComplianceResponse(
        rules=[ComplianceRule(id=1, content="规则")],
        report=ComplianceReport(
            total_rules=1, total_segments_checked=1, violations=violations
        ),
        processing_time_ms=1.0,
    )
    return resp.model_dump(mode="json")


@pytest.fixture
def api(mock_task_store: MagicMock) -> TestClient:
    app = FastAPI()
    app.state.task_store = mock_task_store
    app.include_router(compliance_router)
    return TestClient(app)


class TestViolationIds:
    def test_ids_assigned_and_unique_even_for_same_time_and_rule(self):
        # 旧的 violationKey = timestamp_ms-rule_id 在此场景会冲突
        report = ComplianceReport(
            total_rules=1,
            total_segments_checked=1,
            violations=[_violation(1, 1000), _violation(1, 1000)],
        )
        ids = [v.id for v in report.violations]
        assert all(ids)
        assert len(set(ids)) == 2

    def test_existing_ids_preserved(self):
        report = ComplianceReport(
            total_rules=1,
            total_segments_checked=1,
            violations=[_violation(1, 1000, id="custom")],
        )
        assert report.violations[0].id == "custom"

    def test_legacy_data_without_ids_gets_deterministic_ids(self):
        raw = _response([_violation(1, 1), _violation(2, 2)])
        for v in raw["report"]["violations"]:
            v.pop("id")
        first = ComplianceResponse.model_validate(raw)
        second = ComplianceResponse.model_validate(raw)
        assert [v.id for v in first.report.violations] == [
            v.id for v in second.report.violations
        ]


class TestUpdateViolationStatuses:
    def _setup(self, store: MagicMock, violations: list[Violation]) -> None:
        store.persistence.load_json.return_value = _response(violations)

    def _saved(self, store: MagicMock) -> ComplianceResponse:
        saved = store.persistence.save_json.call_args.args[2]
        return saved

    def test_update_by_violation_id(self, api, mock_task_store):
        self._setup(mock_task_store, [_violation(1, 1000), _violation(2, 2000)])

        r = api.patch(
            "/api/v1/tasks/t/compliance/violations",
            json={"updates": [{"violation_id": "v0002", "status": "confirmed"}]},
        )

        assert r.status_code == 200
        assert r.json() == {"ok": True, "updated": 1, "missing": [], "compliance_score": 94.0}
        statuses = [v.status for v in self._saved(mock_task_store).report.violations]
        assert statuses == ["pending", "confirmed"]

    def test_legacy_index_still_supported(self, api, mock_task_store):
        self._setup(mock_task_store, [_violation(1, 1000), _violation(2, 2000)])

        r = api.patch(
            "/api/v1/tasks/t/compliance/violations",
            json={"updates": [{"index": 0, "status": "rejected"}]},
        )

        assert r.json()["updated"] == 1
        statuses = [v.status for v in self._saved(mock_task_store).report.violations]
        assert statuses == ["rejected", "pending"]

    def test_unknown_targets_reported_as_missing(self, api, mock_task_store):
        self._setup(mock_task_store, [_violation(1, 1000)])

        r = api.patch(
            "/api/v1/tasks/t/compliance/violations",
            json={
                "updates": [
                    {"violation_id": "nope", "status": "confirmed"},
                    {"index": 99, "status": "confirmed"},
                ]
            },
        )

        assert r.json() == {
            "ok": True, "updated": 0, "missing": ["nope", "99"], "compliance_score": 97.0
        }

    def test_update_requires_a_target(self, api, mock_task_store):
        self._setup(mock_task_store, [_violation(1, 1000)])

        r = api.patch(
            "/api/v1/tasks/t/compliance/violations",
            json={"updates": [{"status": "confirmed"}]},
        )

        assert r.status_code == 422

    def test_missing_compliance_json_is_404(self, api, mock_task_store):
        mock_task_store.persistence.load_json.return_value = None

        r = api.patch(
            "/api/v1/tasks/t/compliance/violations",
            json={"updates": [{"violation_id": "v0001", "status": "confirmed"}]},
        )

        assert r.status_code == 404


class TestReviewTrailAndScore:
    def _saved(self, store: MagicMock) -> ComplianceResponse:
        return store.persistence.save_json.call_args.args[2]

    def test_confirm_records_time_and_note(self, api, mock_task_store):
        mock_task_store.persistence.load_json.return_value = _response([_violation(1, 1000)])

        api.patch(
            "/api/v1/tasks/t/compliance/violations",
            json={"updates": [{"violation_id": "v0001", "status": "confirmed", "note": " 已核实原文 "}]},
        )

        v = self._saved(mock_task_store).report.violations[0]
        assert v.status == "confirmed"
        assert v.reviewed_at is not None
        assert v.review_note == "已核实原文"

    def test_back_to_pending_clears_trail(self, api, mock_task_store):
        reviewed = _violation(
            1, 1000, status="rejected", reviewed_at="2026-01-01T00:00:00+00:00", review_note="误报"
        )
        mock_task_store.persistence.load_json.return_value = _response([reviewed])

        api.patch(
            "/api/v1/tasks/t/compliance/violations",
            json={"updates": [{"violation_id": "v0001", "status": "pending"}]},
        )

        v = self._saved(mock_task_store).report.violations[0]
        assert v.reviewed_at is None and v.review_note is None

    def test_rejected_violation_no_longer_deducts_score(self, api, mock_task_store):
        mock_task_store.persistence.load_json.return_value = _response(
            [_violation(1, 1000, severity="high"), _violation(2, 2000, severity="medium")]
        )

        r = api.patch(
            "/api/v1/tasks/t/compliance/violations",
            json={"updates": [{"violation_id": "v0001", "status": "rejected"}]},
        )

        assert r.json()["compliance_score"] == 92.0  # 仅剩 medium 扣 8 分
        assert self._saved(mock_task_store).report.compliance_score == 92.0

    def test_note_length_limited(self, api, mock_task_store):
        mock_task_store.persistence.load_json.return_value = _response([_violation(1, 1000)])
        r = api.patch(
            "/api/v1/tasks/t/compliance/violations",
            json={"updates": [{"violation_id": "v0001", "status": "rejected", "note": "x" * 501}]},
        )
        assert r.status_code == 422


class TestExport:
    def test_export_xlsx_contains_sheets_and_review_info(self, api, mock_task_store):
        import io

        from openpyxl import load_workbook

        mock_task_store.persistence.load_json.return_value = _response(
            [_violation(1, 1000, status="confirmed", review_note="属实", original_text="原文")]
        )

        r = api.get("/api/v1/tasks/abcdef123456/compliance/export")

        assert r.status_code == 200
        assert "compliance_abcdef12.xlsx" in r.headers["content-disposition"]
        wb = load_workbook(io.BytesIO(r.content))
        assert wb.sheetnames == ["概览", "违规明细"]
        detail = wb["违规明细"]
        header = [c.value for c in detail[1]]
        row = dict(zip(header, [c.value for c in detail[2]]))
        assert row["复核状态"] == "已确认"
        assert row["复核备注"] == "属实"
        assert row["原文"] == "原文"

    def test_export_neutralizes_formula_injection(self, api, mock_task_store):
        import io

        from openpyxl import load_workbook

        mock_task_store.persistence.load_json.return_value = _response(
            [_violation(1, 1000, original_text="=HYPERLINK(\"http://evil\")")]
        )

        r = api.get("/api/v1/tasks/t/compliance/export")

        cell = load_workbook(io.BytesIO(r.content))["违规明细"]["G2"]
        assert cell.data_type != "f"
        assert cell.value.startswith("=HYPERLINK")

    def test_export_missing_report_404(self, api, mock_task_store):
        mock_task_store.persistence.load_json.return_value = None
        assert api.get("/api/v1/tasks/t/compliance/export").status_code == 404
